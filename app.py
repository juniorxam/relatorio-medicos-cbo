import csv
import io
import re
import unicodedata
from datetime import datetime
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Relatório de Médicos por Setor", page_icon="logo.png", layout="wide")
st.markdown('''<style>
:root{--navy:#16324F;--teal:#0F766E;--ink:#172033}
.stApp{background:linear-gradient(180deg,#F8FAFC 0%,#FFF 38%);color:var(--ink)}
[data-testid="stSidebar"]{background:linear-gradient(180deg,#102A43,#16324F)}
[data-testid="stSidebar"] *{color:#F8FAFC!important}
[data-testid="stSidebar"] [data-testid="stFileUploader"] section{background:#FFF!important;border:1px solid #CBD5E1!important}
[data-testid="stSidebar"] [data-testid="stFileUploader"] section div{color:#16324F!important}
[data-testid="stSidebar"] [data-testid="stFileUploader"] button{color:#16324F!important;background:#F8FAFC!important;border:1px solid #94A3B8!important}
[data-testid="stSidebar"] [data-testid="stFileUploader"] button *{color:#16324F!important}
[data-testid="stSidebar"] [data-baseweb="select"]>div{background:#FFF!important;border-color:#94A3B8!important}
[data-testid="stSidebar"] [data-baseweb="select"] *,[data-testid="stSidebar"] input{color:#16324F!important;-webkit-text-fill-color:#16324F!important}
.hero{padding:26px 30px;border-radius:20px;margin:4px 0 24px;color:#FFF;background:radial-gradient(circle at 90% 10%,#2DD4BF 0,transparent 30%),linear-gradient(120deg,#16324F,#0F766E);box-shadow:0 12px 28px #0f766e2e}
.hero h1{margin:0;font-size:2rem}.hero p{margin:8px 0 0;color:#D9FDF5}
div[data-testid="stMetric"]{background:#FFF;border:1px solid #E2E8F0;border-radius:14px;padding:12px;box-shadow:0 5px 16px #0f172a0d}
div[data-testid="stMetricValue"]{color:#16324F}.stButton button,.stDownloadButton button{border-radius:10px;font-weight:650}
</style>''', unsafe_allow_html=True)

EXTENSIONS=["csv","xls","xlsx","xlsm"]
ALIASES={
 "numero funcional":"NUMFUNC","numero do funcionario":"NUMFUNC","numero vinculo":"NUMVINC",
 "nome servidor":"SERVIDOR","descricao escala":"DESC. ESCALA","desc escala":"DESC. ESCALA",
 "setor":"SETOR","cargo":"CARGO","ocupacao":"OCUPACAO","cpf":"CPF",
 "carga horaria escalada":"CARGA HORARIA ESCALADA","carga horaria":"CARGA HORARIA","escala":"ESCALA",
}
REQUIRED=["NUMFUNC","NUMVINC","SETOR","CARGO"]
REPORT_COLUMNS=["ORDEM","ESCALA","DESC. ESCALA","NUMFUNC","NUMVINC","SERVIDOR","CPF","SETOR","CARGO","OCUPACAO","CARGA HORARIA ESCALADA","CARGA HORARIA"]
NUMERIC_REPORT_COLUMNS={"ORDEM","ESCALA","NUMFUNC","NUMVINC"}
UNIQUE_SERVER_COLUMNS=["ORDEM","NUMFUNC","NUMVINC","SERVIDOR","SETOR","CARGO","OCUPACAO","CARGA HORARIA"]


def text(v):
 if v is None:return ""
 try:
  if pd.isna(v):return ""
 except (TypeError,ValueError):pass
 s=str(v).strip(); return "" if s.casefold() in {"nan","none","nat"} else s

def key(v):
 s="".join(c for c in unicodedata.normalize("NFKD",text(v).casefold()) if not unicodedata.combining(c))
 return re.sub(r"[^a-z0-9]+"," ",s).strip()

def norm(v): return key(v).upper()

def identificador(v):
 s=text(v).replace(" ","")
 # Excel pode entregar identificadores inteiros como texto decimal (ex.: 123.0).
 if re.fullmatch(r"\d+[\.,]0+",s):s=re.split(r"[\.,]",s,1)[0]
 digits=re.sub(r"\D","",s)
 return digits.lstrip("0") or "0" if digits else ""

def chaves_vinculo(df):
 keys=[]
 for index,row in df.reset_index(drop=True).iterrows():
  numfunc=identificador(row.get("NUMFUNC",""));numvinc=identificador(row.get("NUMVINC",""))
  keys.append(f"{numfunc}-{numvinc}" if numfunc and numvinc else f"SEM_CHAVE_{index+1}")
 return pd.Series(keys,index=df.index,dtype="string")

def chaves_medico(df):
 keys=[]
 for index,row in df.reset_index(drop=True).iterrows():
  numfunc=identificador(row.get("NUMFUNC",""))
  if numfunc:
   keys.append(f"NUMFUNC_{numfunc}");continue
  cpf=identificador(row.get("CPF",""))
  keys.append(f"CPF_{cpf}" if cpf else f"SEM_IDENTIFICADOR_{index+1}")
 return pd.Series(keys,index=df.index,dtype="string")

def medicos_unicos_count(df):
 if df.empty:return 0
 return int(chaves_medico(df).nunique())

def read_raw(data,name):
 if name.casefold().endswith(".csv"):
  for enc in ("utf-8-sig","utf-8","cp1252","latin1"):
   try: raw=data.decode(enc);break
   except UnicodeDecodeError: continue
  try: sep=csv.Sniffer().sniff(raw[:10000],delimiters="\t;,|").delimiter
  except csv.Error: sep=";"
  return pd.read_csv(io.StringIO(raw),sep=sep,header=None,dtype=object,keep_default_na=False)
 errors=[]
 extension=name.casefold().rsplit(".",1)[-1]
 preferred=(['xlrd','openpyxl'] if extension == "xls" or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" else ['openpyxl','xlrd'])
 for engine in preferred:
  try:
   return pd.read_excel(BytesIO(data),engine=engine,header=None,dtype=object)
  except Exception as exc:
   errors.append(f"{engine}: {exc}")
 # Alguns arquivos são CSV/TXT exportados com extensão .xls/.xlsx.
 for enc in ("utf-8-sig","utf-8","cp1252","latin1"):
  try:
   raw=data.decode(enc)
   try: sep=csv.Sniffer().sniff(raw[:10000],delimiters="\t;,|").delimiter
   except csv.Error: sep=";"
   parsed=pd.read_csv(io.StringIO(raw),sep=sep,header=None,dtype=object,keep_default_na=False)
   if parsed.shape[1] > 1:
    return parsed
  except Exception as exc:
   errors.append(f"texto/{enc}: {exc}")
 raise ValueError("Formato de planilha não reconhecido. Tente salvar novamente como XLSX ou CSV. " + " | ".join(errors[:2]))

def parse(raw):
 raw=raw.dropna(how="all").reset_index(drop=True)
 if raw.empty:return pd.DataFrame()
 idx=0
 for i in range(min(20,len(raw))):
  vals={key(x).replace(" ","").upper() for x in raw.iloc[i].tolist()}
  if "CARGO" in vals and "SETOR" in vals:idx=i;break
 body=raw.iloc[idx+1:].copy(); cols=[]; used=set()
 for value in raw.iloc[idx].tolist():
  label=text(value); canonical=ALIASES.get(key(label),label.upper())
  if canonical in used:
   n=2
   while f"{canonical}_{n}" in used:n+=1
   canonical=f"{canonical}_{n}"
  used.add(canonical);cols.append(canonical)
 body.columns=cols; return body.reset_index(drop=True)

def clean(df):
 df=df.copy()
 for c in df.columns: df[c]=df[c].map(text)
 for c in ["NUMFUNC","NUMVINC"]:
  if c not in df:df[c]=""
  df[c]=df[c].map(identificador)
 for c in REQUIRED:
  if c not in df:df[c]=""
 if "VINCULO" not in df:df["VINCULO"]=df.get("TIPO_VINCULO","")
 df["CARGO_NORM"]=df["CARGO"].map(norm)
 df["SETOR"]=df["SETOR"].map(text)
 return df[df["CARGO_NORM"].isin({"MEDICO","MEDICO RQE","MEDICO CLINICO"})].reset_index(drop=True)

def vinculo_categoria(value):
 n=norm(value)
 if any(term in n for term in ("REQUISITADO","REQUISITADA","CEDIDO","CEDIDA")):return "REQUISITADO"
 if any(term in n for term in ("CONTRATO","TEMPORARIO","TEMPORARIA")):return "CONTRATO TEMPORARIO"
 if any(term in n for term in ("CONCURSADO","CONCURSADA","EFETIVO","EFETIVA","ESTATUTARIO")):return "EFETIVO"
 return "NAO CLASSIFICADO"

def cpf_formatado(value):
 digits=re.sub(r"\D","",text(value))
 if not digits:return ""
 digits=digits[-11:].zfill(11)
 return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"

def inteiro_coluna(series):
 values=pd.to_numeric(series.replace("",pd.NA),errors="coerce")
 return values.astype("Int64")

def reference_summary(df):
 columns=["TIPO_LINHA","SETOR","ESPECIALIDADE","EFETIVOS","CONTRATO_TEMPORARIO","REQUISITADO","TOTAL_GERAL"]
 if df.empty:return pd.DataFrame(columns=columns)
 work=df.copy();work["CATEGORIA_VINCULO"]=work["VINCULO"].map(vinculo_categoria);work["ESPECIALIDADE"]=work["OCUPACAO"].replace("","NÃO POSSUI/NÃO CADASTRADA")
 rows=[]
 for setor,sector_df in work.groupby("SETOR",sort=True,dropna=False):
  counts=sector_df["CATEGORIA_VINCULO"].value_counts(); rows.append({"TIPO_LINHA":"SETOR","SETOR":setor,"ESPECIALIDADE":setor,"EFETIVOS":int(counts.get("EFETIVO",0)),"CONTRATO_TEMPORARIO":int(counts.get("CONTRATO TEMPORARIO",0)),"REQUISITADO":int(counts.get("REQUISITADO",0)),"TOTAL_GERAL":len(sector_df)})
  for specialty,specialty_df in sector_df.groupby("ESPECIALIDADE",sort=True,dropna=False):
   counts=specialty_df["CATEGORIA_VINCULO"].value_counts();rows.append({"TIPO_LINHA":"ESPECIALIDADE","SETOR":setor,"ESPECIALIDADE":specialty,"EFETIVOS":int(counts.get("EFETIVO",0)),"CONTRATO_TEMPORARIO":int(counts.get("CONTRATO TEMPORARIO",0)),"REQUISITADO":int(counts.get("REQUISITADO",0)),"TOTAL_GERAL":len(specialty_df)})
 counts=work["CATEGORIA_VINCULO"].value_counts()
 rows.append({"TIPO_LINHA":"TOTAL","SETOR":"TOTAL GERAL","ESPECIALIDADE":"TOTAL GERAL","EFETIVOS":int(counts.get("EFETIVO",0)),"CONTRATO_TEMPORARIO":int(counts.get("CONTRATO TEMPORARIO",0)),"REQUISITADO":int(counts.get("REQUISITADO",0)),"TOTAL_GERAL":len(work)})
 return pd.DataFrame(rows,columns=columns)

def report(df):
 out=df.copy()
 for c in REPORT_COLUMNS:
  if c not in out:out[c]=""
 out=out[REPORT_COLUMNS].reset_index(drop=True);out["ORDEM"]=range(1,len(out)+1)
 for c in NUMERIC_REPORT_COLUMNS:out[c]=inteiro_coluna(out[c])
 out["CPF"]=out["CPF"].map(cpf_formatado)
 return out

def unique_servers_report(df):
 columns=UNIQUE_SERVER_COLUMNS
 if df.empty:return pd.DataFrame(columns=columns)
 out=df.copy()
 for c in columns:
  if c not in out:out[c]=""
 # A primeira ocorrência representa o servidor/vínculo; as demais escalas são removidas.
 out["_CHAVE_SERVIDOR"] = chaves_vinculo(out)
 out=out.drop_duplicates("_CHAVE_SERVIDOR",keep="first").drop(columns="_CHAVE_SERVIDOR")[columns].reset_index(drop=True)
 out["ORDEM"]=range(1,len(out)+1)
 for c in {"ORDEM","NUMFUNC","NUMVINC"}:out[c]=inteiro_coluna(out[c])
 return out

def summary_sector(df):
 if df.empty:return pd.DataFrame(columns=["SETOR","MEDICOS","VINCULOS","ESPECIALIDADES"])
 out=df.assign(_CHAVE_MEDICO=chaves_medico(df),_ESP=df["OCUPACAO"].replace("","Não informada"))
 return out.groupby("SETOR",dropna=False).agg(MEDICOS=("_CHAVE_MEDICO","nunique"),VINCULOS=("NUMVINC","count"),ESPECIALIDADES=("_ESP",lambda x:"; ".join(sorted(set(x))))).reset_index().sort_values("SETOR")

def summary_specialty(df):
 if df.empty:return pd.DataFrame(columns=["SETOR","ESPECIALIDADE","MEDICOS"])
 out=df.assign(_CHAVE_MEDICO=chaves_medico(df),ESPECIALIDADE=df["OCUPACAO"].replace("","Não informada"))
 return out.groupby(["SETOR","ESPECIALIDADE"],dropna=False).agg(MEDICOS=("_CHAVE_MEDICO","nunique")).reset_index().sort_values(["SETOR","ESPECIALIDADE"])

def management_base(df):
 columns=["CHAVE_VINCULO","CHAVE_MEDICO","NUMFUNC","NUMVINC","SERVIDOR","SETOR","OCUPACAO","VINCULO","ESCALAS","HORAS_ESCALADAS","CARGA_HORARIA"]
 if df.empty:return pd.DataFrame(columns=columns)
 work=df.copy()
 for c in ["SERVIDOR","SETOR","OCUPACAO","VINCULO"]:
  if c not in work:work[c]=""
 work["CHAVE_VINCULO"]=chaves_vinculo(work)
 work["CHAVE_MEDICO"]=chaves_medico(work)
 work["HORAS_ESCALADAS_NUM"]=pd.to_numeric(work.get("CARGA HORARIA ESCALADA",pd.Series(index=work.index)),errors="coerce").fillna(0)
 work["CARGA_HORARIA_NUM"]=pd.to_numeric(work.get("CARGA HORARIA",pd.Series(index=work.index)),errors="coerce").fillna(0)
 out=work.groupby("CHAVE_VINCULO",as_index=False).agg(CHAVE_MEDICO=("CHAVE_MEDICO","first"),NUMFUNC=("NUMFUNC","first"),NUMVINC=("NUMVINC","first"),SERVIDOR=("SERVIDOR","first"),SETOR=("SETOR","first"),OCUPACAO=("OCUPACAO","first"),VINCULO=("VINCULO","first"),ESCALAS=("CHAVE_VINCULO","size"),HORAS_ESCALADAS=("HORAS_ESCALADAS_NUM","sum"),CARGA_HORARIA=("CARGA_HORARIA_NUM","max"))
 return out[columns]

def management_sector(df):
 base=management_base(df)
 columns=["SETOR","MEDICOS","ESCALAS","HORAS_ESCALADAS","CARGA_HORARIA","SALDO_CARGA","COBERTURA_PCT"]
 if base.empty:return pd.DataFrame(columns=columns)
 out=base.groupby("SETOR",dropna=False).agg(MEDICOS=("CHAVE_MEDICO","nunique"),ESCALAS=("ESCALAS","sum"),HORAS_ESCALADAS=("HORAS_ESCALADAS","sum"),CARGA_HORARIA=("CARGA_HORARIA","sum")).reset_index()
 out["SALDO_CARGA"]=out["CARGA_HORARIA"]-out["HORAS_ESCALADAS"]
 out["COBERTURA_PCT"]=out.apply(lambda r:round(r.HORAS_ESCALADAS/r.CARGA_HORARIA*100,1) if r.CARGA_HORARIA else 0,axis=1)
 return out.sort_values("SETOR")

def management_specialty(df):
 base=management_base(df)
 if base.empty:return pd.DataFrame(columns=["OCUPACAO","MEDICOS","ESCALAS","HORAS_ESCALADAS"])
 return base.assign(OCUPACAO=base["OCUPACAO"].replace("","Não informada")).groupby("OCUPACAO",dropna=False).agg(MEDICOS=("CHAVE_MEDICO","nunique"),ESCALAS=("ESCALAS","sum"),HORAS_ESCALADAS=("HORAS_ESCALADAS","sum")).reset_index().sort_values("MEDICOS",ascending=False)

def management_vinculo(df):
 base=management_base(df)
 if base.empty:return pd.DataFrame(columns=["VINCULO","MEDICOS","ESCALAS"])
 return base.assign(VINCULO=base["VINCULO"].replace("","Não informado")).groupby("VINCULO",dropna=False).agg(MEDICOS=("CHAVE_MEDICO","nunique"),ESCALAS=("ESCALAS","sum")).reset_index().sort_values("MEDICOS",ascending=False)

def quality_summary(df):
 if df.empty:return pd.DataFrame(columns=["INDICADOR","REGISTROS"])
 checks={"Sem CPF":df.get("CPF",pd.Series(index=df.index)).map(text).eq("").sum(),"Sem setor":df.get("SETOR",pd.Series(index=df.index)).map(text).eq("").sum(),"Sem especialidade/ocupação":df.get("OCUPACAO",pd.Series(index=df.index)).map(text).eq("").sum(),"Sem carga escalada":pd.to_numeric(df.get("CARGA HORARIA ESCALADA",pd.Series(index=df.index)),errors="coerce").isna().sum(),"Vínculo não classificado":df.get("VINCULO",pd.Series(index=df.index)).map(vinculo_categoria).eq("NAO CLASSIFICADO").sum()}
 return pd.DataFrame([{"INDICADOR":k,"REGISTROS":int(v)} for k,v in checks.items()])

def style(ws,df):
 fill=PatternFill("solid",fgColor="16324F");font=Font(bold=True,color="FFFFFF");side=Side(style="thin",color="CBD5E1")
 for i,c in enumerate(df.columns,1):
  cell=ws.cell(1,i);cell.fill=fill;cell.font=font;cell.alignment=Alignment(horizontal="center",wrap_text=True);cell.border=Border(left=side,right=side,top=side,bottom=side)
  if c in NUMERIC_REPORT_COLUMNS:ws.column_dimensions[get_column_letter(i)].number_format="0"
  vals=[len(str(c))]+([int(df[c].astype("string").str.len().fillna(0).max())] if len(df) else []);ws.column_dimensions[get_column_letter(i)].width=min(max(max(vals)+2,12),50)
 for row in ws.iter_rows(min_row=2,max_row=ws.max_row,max_col=len(df.columns)):
  for cell in row:
   cell.border=Border(left=side,right=side,top=side,bottom=side);cell.alignment=Alignment(vertical="center")
   if df.columns[cell.column-1] in NUMERIC_REPORT_COLUMNS:cell.number_format="0";cell.alignment=Alignment(horizontal="right",vertical="center")
 if len(df):ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;ws.print_title_rows="1:1";ws.print_area=f"A1:{get_column_letter(len(df.columns))}{ws.max_row}"
 ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True

def excel(sheets):
 out=BytesIO()
 with pd.ExcelWriter(out,engine="openpyxl") as writer:
  for name,df in sheets.items():df.to_excel(writer,sheet_name=name[:31],index=False);style(writer.sheets[name[:31]],df)
 return out.getvalue()

def pdf(reference):
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A4,landscape
 from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
 from reportlab.lib.units import mm
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
 from xml.sax.saxutils import escape
 out=BytesIO();doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=10*mm,rightMargin=10*mm,topMargin=9*mm,bottomMargin=9*mm,title="Número de médicos por especialidade")
 styles=getSampleStyleSheet();title=ParagraphStyle("t",parent=styles["Title"],fontSize=16,textColor=colors.HexColor("#16324F"),alignment=1,spaceAfter=3);sub=ParagraphStyle("s",parent=styles["Normal"],fontSize=8,textColor=colors.HexColor("#64748B"),alignment=1,spaceAfter=7);cell=ParagraphStyle("c",parent=styles["Normal"],fontSize=6.8,leading=7.8);head=ParagraphStyle("h",parent=cell,textColor=colors.white,fontName="Helvetica-Bold",alignment=1)
 story=[Paragraph("NÚMERO DE MÉDICOS POR ESPECIALIDADE",title),Paragraph(f"Gerado em {datetime.today().strftime('%d/%m/%Y %H:%M')} | Cargos: MÉDICO, MEDICO - RQE e MEDICO CLINICO",sub)]
 data=[[Paragraph("LOTAÇÃO/ESPECIALIDADE",head),Paragraph("Efetivos",head),Paragraph("Contrato Temporário",head),Paragraph("Requisitado",head),Paragraph("Total Geral",head)]]
 for _,r in reference.iterrows():
  label=text(r.ESPECIALIDADE); style=ParagraphStyle("row",parent=cell,fontName="Helvetica-Bold" if r.TIPO_LINHA=="SETOR" else "Helvetica",leftIndent=0 if r.TIPO_LINHA=="SETOR" else 12)
  data.append([Paragraph(escape(label),style),str(r.EFETIVOS or "-"),str(r.CONTRATO_TEMPORARIO or "-"),str(r.REQUISITADO or "-"),str(r.TOTAL_GERAL or "-")])
 if len(data)==1:data.append([Paragraph("Nenhum médico encontrado.",cell),"-","-","-","-"])
 table=Table(data,colWidths=[155*mm,27*mm,38*mm,27*mm,27*mm],repeatRows=1)
 table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#16324F")),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#CBD5E1")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8FAFC")]),("BACKGROUND",(0,1),(-1,-1),colors.white),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(1,1),(-1,-1),"CENTER")]))
 for index,row in enumerate(reference.itertuples(),start=1):
  if row.TIPO_LINHA=="SETOR":table.setStyle(TableStyle([("BACKGROUND",(0,index),(-1,index),colors.HexColor("#E6FFFA")),("FONTNAME",(0,index),(-1,index),"Helvetica-Bold")]))
  if row.TIPO_LINHA=="TOTAL":table.setStyle(TableStyle([("BACKGROUND",(0,index),(-1,index),colors.HexColor("#D1FAE5")),("FONTNAME",(0,index),(-1,index),"Helvetica-Bold"),("LINEABOVE",(0,index),(-1,index),1,colors.HexColor("#0F766E"))]))
 story.append(table);doc.build(story);return out.getvalue()

def main():
 st.markdown('<div class="hero"><div>RELATÓRIO DE CBO</div><h1>Médicos por setor</h1><p>Relação de médicos e especialidades cadastradas em cada setor.</p></div>',unsafe_allow_html=True)
 with st.sidebar:
  st.header("Arquivo da Intranet");upload=st.file_uploader("Planilha de escalas",type=EXTENSIONS);st.caption("São considerados somente os cargos MÉDICO, MEDICO - RQE e MEDICO CLINICO.")
 if not upload:st.info("Envie a planilha da Intranet para gerar o relatório.");st.stop()
 try:
  raw=parse(read_raw(upload.getvalue(),upload.name));missing=[c for c in REQUIRED if c not in raw.columns]
  if missing:st.error("Colunas obrigatórias ausentes: "+", ".join(missing));st.stop()
  doctors=clean(raw)
 except Exception as e:st.error(f"Não foi possível processar a planilha: {e}");st.stop()
 sectors=["(Todos)"]+sorted(doctors.SETOR.unique().tolist()) if len(doctors) else ["(Todos)"]
 vinculos=["(Todos)"]+sorted([text(v) for v in doctors["VINCULO"].unique() if text(v)]) if len(doctors) else ["(Todos)"]
 especialidades=["(Todas)"]+sorted([text(v) for v in doctors["OCUPACAO"].unique() if text(v)]) if len(doctors) else ["(Todas)"]
 with st.sidebar:
  sector=st.selectbox("Filtrar setor",sectors)
  vinculo=st.selectbox("Filtrar vínculo",vinculos)
  especialidade=st.selectbox("Filtrar especialidade",especialidades)
  unique_enabled=st.checkbox("Gerar relatório de servidores sem repetição",value=True,help="Remove registros repetidos usando NUMFUNC + NUMVINC como chave.")
 filtered=doctors if sector=="(Todos)" else doctors[doctors.SETOR==sector]
 if vinculo!="(Todos)":filtered=filtered[filtered["VINCULO"].map(text)==vinculo]
 if especialidade!="(Todas)":filtered=filtered[filtered["OCUPACAO"].map(text)==especialidade]
 detail=report(filtered);unique_servers=unique_servers_report(filtered);sector_summary=summary_sector(doctors);specialty_summary=summary_specialty(doctors);reference=reference_summary(doctors)
 base_management=management_base(filtered);sector_management=management_sector(filtered);specialty_management=management_specialty(filtered);vinculo_management=management_vinculo(filtered);quality=quality_summary(filtered)
 unique_count=medicos_unicos_count(filtered);total_scales=len(filtered);total_hours=base_management["HORAS_ESCALADAS"].sum() if len(base_management) else 0
 missing_key=((filtered["NUMFUNC"].map(identificador)=="")|(filtered["NUMVINC"].map(identificador)=="")).sum()
 k1,k2,k3,k4=st.columns(4);k1.metric("Médicos únicos",unique_count);k2.metric("Escalas/registros",total_scales);k3.metric("Setores",filtered["SETOR"].nunique());k4.metric("Horas escaladas",f"{total_hours:,.0f}".replace(",","."))
 if missing_key:st.warning(f"{missing_key} registro(s) não possuem NUMFUNC ou NUMVINC. Eles não são agrupados entre si, pois não é possível confirmar que representam o mesmo vínculo.")
 tab_names=["Painel gerencial","Capacidade e carga","Qualidade dos dados","Relação de médicos","Resumo no padrão do PDF","Especialidades por setor"]
 if unique_enabled:tab_names.append("Servidores sem repetição")
 tabs=st.tabs(tab_names)
 with tabs[0]:
  st.subheader("Visão geral para gerenciamento")
  a,b=st.columns(2)
  with a:
   st.markdown("**Médicos por setor**")
   if len(sector_management):st.bar_chart(sector_management.set_index("SETOR")["MEDICOS"])
  with b:
   st.markdown("**Médicos por especialidade**")
   if len(specialty_management):st.bar_chart(specialty_management.set_index("OCUPACAO")["MEDICOS"].head(15))
  st.markdown("**Distribuição por vínculo**");st.dataframe(vinculo_management,use_container_width=True,hide_index=True)
 with tabs[1]:
  st.subheader("Capacidade e carga horária")
  st.caption("A carga escalada é somada por escala; a carga horária contratual usa o maior valor informado por NUMFUNC + NUMVINC para evitar duplicação.")
  st.dataframe(sector_management,use_container_width=True,hide_index=True)
  if len(sector_management):st.bar_chart(sector_management.set_index("SETOR")[["HORAS_ESCALADAS","CARGA_HORARIA"]])
 with tabs[2]:
  st.subheader("Qualidade cadastral")
  st.dataframe(quality,use_container_width=True,hide_index=True)
  st.caption("Use estes alertas para priorizar correções no cadastro da Intranet.")
 with tabs[3]:
  if detail.empty:st.warning("Nenhum médico encontrado com os três cargos definidos.")
  else:st.dataframe(detail,use_container_width=True,hide_index=True,height=520)
 with tabs[4]:st.dataframe(reference.drop(columns=["TIPO_LINHA"]),use_container_width=True,hide_index=True)
 with tabs[5]:st.dataframe(specialty_summary,use_container_width=True,hide_index=True)
 if unique_enabled:
  with tabs[6]:
   removed=max(len(filtered)-len(unique_servers),0)
   st.success(f"{len(unique_servers)} servidores únicos. {removed} registro(s) repetido(s) removido(s) pela chave NUMFUNC + NUMVINC.")
   st.dataframe(unique_servers,use_container_width=True,hide_index=True)
 st.markdown("---");st.subheader("Exportar relatórios")
 sheets={"Medicos":detail,"Resumo setores":sector_summary,"Especialidades":specialty_summary,"Resumo padrão PDF":reference}
 if unique_enabled:sheets["Servidores unicos"]=unique_servers
 xlsx=excel(sheets);p=pdf(reference)
 a,b=st.columns(2)
 with a:st.download_button("Baixar Excel",xlsx,f"relatorio_medicos_cbo_{datetime.today():%Y%m%d}.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
 with b:st.download_button("Baixar resumo em PDF",p,f"resumo_medicos_cbo_{datetime.today():%Y%m%d}.pdf","application/pdf",use_container_width=True)
if __name__=="__main__":main()
